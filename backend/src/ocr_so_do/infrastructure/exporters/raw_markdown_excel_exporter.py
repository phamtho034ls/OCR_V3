"""
infrastructure/exporters/raw_markdown_excel_exporter.py
Chuyen doi du lieu tho dang Markdown sang cau truc bang tinh Excel (.xlsx)
ho tro hien thi xem truoc (preview) tren Web UI va tai ve tep Excel hoan chinh.
"""
import io
import re
import logging
import unicodedata
from typing import Dict, Any, List, Optional, Tuple
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .excel_129_exporter import Excel129Exporter
from ...application.projections.cadastral_129_mapper import Cadastral129Mapper
from ...domain.rules.validation.validators import GCNValidators

logger = logging.getLogger(__name__)


class RawMarkdownExcelExporter:
    """
    Phan tich van ban Markdown tho va xuat sang dinh dang Excel (.xlsx).
    """

    @staticmethod
    def parse_raw_markdown(raw_markdown: str) -> Dict[str, Any]:
        """
        Phan ra Markdown tho thanh cau truc du lieu gom:
        - metadata: {job_id, file_name, template, total_pages, created_at}
        - summary_fields: List[Dict[str, Any]] (stt, field, value)
        - ocr_lines: List[Dict[str, Any]] (stt, page, file_name, text)
        - merged_dict: Dict[str, Any]
        """
        result: Dict[str, Any] = {
            "metadata": {},
            "summary_fields": [],
            "ocr_lines": [],
            "merged_dict": {}
        }
        if not raw_markdown:
            return result

        # 1. Trich xuat metadata tu block header
        # Pattern: > **Ma Ho So / Job ID:** `...` | **Ten file:** `...` | ...
        for line in raw_markdown.splitlines():
            line_s = line.strip()
            if line_s.startswith(">") and "Job ID" in line_s:
                parts = line_s.lstrip(">").split("|")
                for p in parts:
                    if "**Mã Hồ Sơ" in p or "**Job ID" in p:
                        m = re.search(r"`([^`]+)`", p)
                        if m:
                            result["metadata"]["job_id"] = m.group(1).strip()
                    elif "**Tên file" in p or "**File" in p:
                        m = re.search(r"`([^`]+)`", p)
                        if m:
                            result["metadata"]["file_name"] = m.group(1).strip()
                    elif "**Mẫu Sổ" in p or "**Template" in p:
                        m = re.search(r"`([^`]+)`", p)
                        if m:
                            result["metadata"]["template"] = m.group(1).strip()
                    elif "**Tổng số trang" in p or "**Pages" in p:
                        m = re.search(r"`([^`]+)`", p)
                        if m:
                            result["metadata"]["total_pages"] = m.group(1).strip()
                    elif "**Thời điểm" in p or "**Time" in p:
                        m = re.search(r"`([^`]+)`", p)
                        if m:
                            result["metadata"]["created_at"] = m.group(1).strip()
                break

        merged_dict: Dict[str, Any] = {
            "nguoi_su_dung": {},
            "thua_dat": {},
            "cap_gcn": {},
            "bien_dong": {}
        }
        result["merged_dict"] = merged_dict

        sec1_match = re.search(r"## I\. DỮ LIỆU BÓC TÁCH THEO LOGIC.*?\n```(?:text)?\n(.*?)\n```", raw_markdown, re.DOTALL)
        if sec1_match:
            lines = sec1_match.group(1).split("\n")
            stt = 1
            in_danh_sach = False
            danh_sach_thua = []
            for line in lines:
                line = line.strip()
                if not line:
                    continue
                if "[DANH SÁCH CHI TIẾT CÁC THỬA ĐẤT]" in line:
                    in_danh_sach = True
                    continue

                if in_danh_sach:
                    m_head = re.match(r"Thửa\s*(\d+)\s*:\s*Thửa số\s*([^|]+)", line, re.IGNORECASE)
                    if m_head:
                        stt_val = int(m_head.group(1))
                        parts = [p.strip() for p in line.split("|")]
                        s_thua = m_head.group(2).strip()
                        if s_thua == "-": s_thua = ""
                        values: Dict[str, str] = {}
                        for part in parts[1:]:
                            m_map = re.match(r"^Tờ số\s+(.+)$", part, re.IGNORECASE)
                            if m_map:
                                values["tờ số"] = m_map.group(1).strip()
                                continue
                            if ":" not in part:
                                continue
                            key, value = part.split(":", 1)
                            values[key.strip().lower()] = value.strip()
                        t_bando = values.get("tờ số", "")
                        if t_bando == "-": t_bando = ""
                        dt_val = re.sub(r"\s*m(?:2|²)\s*$", "", values.get("diện tích", ""), flags=re.IGNORECASE).strip()
                        if dt_val in ["-", "None"]: dt_val = ""
                        md_val = values.get("mục đích", "")
                        if md_val == "-": md_val = ""
                        th_val = values.get("thời hạn", "")
                        if th_val == "-" or "mục đích" in th_val.lower(): th_val = ""
                        ng_val = values.get("nguồn gốc", "")
                        if ng_val == "-": ng_val = ""
                        dc_val = values.get("địa chỉ", "")
                        if dc_val == "-": dc_val = ""
                        danh_sach_thua.append({
                            "stt": stt_val,
                            "so_thua": s_thua,
                            "to_ban_do": t_bando,
                            "dien_tich": dt_val,
                            "ma_muc_dich": md_val,
                            "muc_dich_su_dung": md_val,
                            "thoi_han": th_val,
                            "nguon_goc": ng_val,
                            "dia_chi": dc_val,
                        })
                    continue

                if not line or ":" not in line:
                    continue
                k, v = line.split(":", 1)
                field_name = k.strip()
                val = v.strip()
                if val == "-":
                    val = ""
                result["summary_fields"].append({
                    "stt": stt,
                    "field": field_name,
                    "value": val
                })
                stt += 1

                fn_low = field_name.lower()
                if "họ và tên chủ 1" in fn_low:
                    merged_dict["nguoi_su_dung"]["ho_ten_chu_1"] = val
                elif "năm sinh chủ 1" in fn_low:
                    merged_dict["nguoi_su_dung"]["ngay_sinh_chu_1"] = val
                elif "số cmnd/cccd chủ 1" in fn_low or "cccd chủ 1" in fn_low:
                    merged_dict["nguoi_su_dung"]["cmnd_chu_1"] = val
                elif "họ và tên chủ 2" in fn_low:
                    merged_dict["nguoi_su_dung"]["ho_ten_chu_2"] = val
                elif "năm sinh chủ 2" in fn_low:
                    merged_dict["nguoi_su_dung"]["ngay_sinh_chu_2"] = val
                elif "số cmnd/cccd chủ 2" in fn_low:
                    merged_dict["nguoi_su_dung"]["cmnd_chu_2"] = val
                elif "địa chỉ thường trú chủ 2" in fn_low:
                    merged_dict["nguoi_su_dung"]["dia_chi_thuong_tru_chu_2"] = val
                elif "địa chỉ thường trú" in fn_low:
                    merged_dict["nguoi_su_dung"]["dia_chi_thuong_tru"] = val
                elif "thửa đất số" in fn_low or "số thửa" in fn_low:
                    merged_dict["thua_dat"]["so_thua"] = val
                elif "tờ bản đồ số" in fn_low:
                    merged_dict["thua_dat"]["to_ban_do"] = val
                elif "địa chỉ thửa đất" in fn_low:
                    merged_dict["thua_dat"]["dia_chi"] = val
                elif "diện tích" in fn_low and "bằng chữ" not in fn_low:
                    m_dt = re.search(r"([\d\.,]+)\s*m2", val)
                    merged_dict["thua_dat"]["dien_tich"] = m_dt.group(1) if m_dt else val
                elif "hình thức sử dụng" in fn_low:
                    merged_dict["thua_dat"]["hinh_thuc_su_dung"] = val
                elif "mục đích sử dụng" in fn_low:
                    merged_dict["thua_dat"]["muc_dich_su_dung"] = val
                elif "thời hạn sử dụng" in fn_low:
                    merged_dict["thua_dat"]["thoi_han"] = val
                elif "nguồn gốc" in fn_low:
                    merged_dict["thua_dat"]["nguon_goc"] = val
                elif "cơ quan cấp" in fn_low:
                    merged_dict["cap_gcn"]["noi_cap"] = val
                elif "ngày cấp gcn" in fn_low:
                    merged_dict["cap_gcn"]["ngay_cap"] = val
                elif "người ký gcn" in fn_low:
                    merged_dict["cap_gcn"]["nguoi_ky_qd"] = re.sub(r"\s*\([^)]*\)\s*$", "", val).strip()
                elif "số vào sổ" in fn_low:
                    merged_dict["so_vao_so"] = val
                elif "số phát hành" in fn_low or "serial" in fn_low:
                    merged_dict["so_phat_hanh"] = val
                elif "mã vạch" in fn_low or "barcode" in fn_low:
                    merged_dict["ma_vach"] = val

            # Lọc danh sách thửa: loại bỏ bản ghi rỗng số thửa
            valid_ds = [p for p in danh_sach_thua if str(p.get("so_thua") or "").strip() and str(p.get("so_thua") or "").strip() != "-"]
            unique_st = set(str(p.get("so_thua") or "").strip() for p in valid_ds)
            # Chỉ coi là nhiều thửa nếu có từ 2 thửa hợp lệ và số thửa khác nhau
            if len(valid_ds) >= 2 and len(unique_st) >= 2:
                # Kế thừa nguồn gốc / tờ bản đồ / thời hạn cấp sổ nếu từng thửa bị thiếu
                general_ng = merged_dict.get("thua_dat", {}).get("nguon_goc", "")
                general_tb = merged_dict.get("thua_dat", {}).get("to_ban_do", "")
                general_th = merged_dict.get("thua_dat", {}).get("thoi_han", "")
                for p_item in valid_ds:
                    if not p_item.get("nguon_goc") and general_ng:
                        p_item["nguon_goc"] = general_ng
                    if not p_item.get("to_ban_do") and general_tb:
                        p_item["to_ban_do"] = general_tb
                    if not p_item.get("thoi_han") and general_th:
                        p_item["thoi_han"] = general_th
                merged_dict["thua_dat"]["danh_sach_thua"] = valid_ds
            else:
                merged_dict["thua_dat"]["danh_sach_thua"] = []
            result["merged_dict"] = merged_dict

        # 3. Trich xuat Section II (Cac trang va dong text OCR)
        page_pattern = re.compile(
            r"###\s*(?:📄|📃)?\s*Trang\s*(\d+)[:\s]*`?([^`\n\(\)]*)`?.*?\n```(?:text)?\n(.*?)\n```",
            re.DOTALL
        )
        page_blocks = page_pattern.findall(raw_markdown)
        line_stt = 1
        if page_blocks:
            for p_num, p_file, p_content in page_blocks:
                page_idx = int(p_num) if p_num.isdigit() else 1
                page_file = p_file.strip() or f"Trang_{page_idx}.png"
                for text_line in p_content.split("\n"):
                    text_line = text_line.strip()
                    if text_line and not text_line.startswith("[") and not text_line.endswith("]"):
                        result["ocr_lines"].append({
                            "stt": line_stt,
                            "page": page_idx,
                            "file_name": page_file,
                            "text": text_line
                        })
                        line_stt += 1
        else:
            generic_code_blocks = re.findall(r"```(?:text)?\n(.*?)\n```", raw_markdown, re.DOTALL)
            for block in generic_code_blocks:
                for text_line in block.split("\n"):
                    if text_line and ":" not in text_line[:25]:
                        result["ocr_lines"].append({
                            "stt": line_stt,
                            "page": 1,
                            "file_name": "N/A",
                            "text": text_line
                        })
                        line_stt += 1

        # 4. Fallback: Nếu Số vào sổ ở Section I bị thiếu hoặc là placeholder '-', quét Section II
        curr_svs = merged_dict.get("so_vao_so", "")
        is_svs_ok, _, _ = GCNValidators.validate_registry_book_number(curr_svs) if curr_svs else (False, "", None)
        if is_svs_ok:
            # Ghi lại giá trị đã chuẩn hóa, ví dụ CHC0124 -> CH00124.
            _, normalized_svs, _ = GCNValidators.validate_registry_book_number(curr_svs)
            merged_dict["so_vao_so"] = normalized_svs or ""
        else:
            # Không để giá trị không hợp lệ ở Section I lọt xuống mapper.
            # Nếu Section II không tìm được candidate đủ tin cậy thì giữ rỗng,
            # thay vì xuất các chuỗi bị cắt như CH0/CH00.
            merged_dict["so_vao_so"] = ""
        if not is_svs_ok:
            for item in result["ocr_lines"]:
                t_line = item.get("text", "")
                if "tiếp nhận" in t_line.lower() or "đơn đề nghị" in t_line.lower():
                    continue
                # Quét theo nhãn
                m_label = re.search(
                    r"(?:(?:Số|[0-9]{1,2}|Sẽ|Sé|Sa|Sô|Số|só|sổ|m6|vô)\s*(?:vào|v[aà]o|v[aà]n)?\s*(?:s[ổốóoôòõỏ]|có)\s*(?:cấp|c[aâấ]p|có)?\s*(?:GCN|GƠN|GTN|sổ)?)\s*[:\.]?\s*([A-Za-z0-9\.\-_/% ]+)",
                    t_line,
                    re.IGNORECASE
                )
                if m_label:
                    cand = m_label.group(1).strip()
                    cand = re.split(r"[\n\r]|ngày|ngay|năm|tháng|bải|số\s*phát\s*hành", cand, flags=re.IGNORECASE)[0].strip()
                    ok_cand, norm_cand, _ = GCNValidators.validate_registry_book_number(cand)
                    if ok_cand and norm_cand:
                        merged_dict["so_vao_so"] = norm_cand
                        is_svs_ok = True
                        break
                # Quét pattern CH/CS
                m_ch = re.search(r"(?:GCN|GƠN|sổ)?\s*(C[HNS]\s*[0-9OoSBDlIqG%T]{3,8}|\d{3,6}\s*/\s*QĐ)", t_line, re.IGNORECASE)
                if m_ch:
                    ok_cand, norm_cand, _ = GCNValidators.validate_registry_book_number(m_ch.group(1))
                    if ok_cand and norm_cand:
                        merged_dict["so_vao_so"] = norm_cand
                        is_svs_ok = True
                        break

        # 5. Fallback: Cơ quan cấp GCN (GCN_donViCap) nếu Section I đang thiếu
        curr_auth = merged_dict.get("cap_gcn", {}).get("noi_cap", "")
        if not curr_auth or curr_auth in ["-", "None"]:
            found_auth = ""
            for item in result["ocr_lines"]:
                t_line = item.get("text", "")
                if "tiếp nhận" in t_line.lower() or "kính gửi" in t_line.lower():
                    continue
                if any(k in t_line.upper() for k in ["UỶ BAN NHÂN DÂN", "ỦY BAN NHÂN DÂN", "UBND", "UÝ BAN"]):
                    val = re.sub(r'^(?:TM\s*\.?\s*)+', '', t_line, flags=re.IGNORECASE).strip(' .:-')
                    val = re.sub(r'^(?:UÝ|UỶ|ỦY)\s*BAN\s*NHÂN\s*DÂN', 'Ủy ban nhân dân', val, flags=re.IGNORECASE)
                    if "lộc bình" in t_line.lower() or "nam quan" in t_line.lower():
                        found_auth = "Ủy ban nhân dân huyện Lộc Bình"
                        break
                    elif len(val) >= 10:
                        found_auth = val
            if not found_auth:
                found_auth = "Ủy ban nhân dân huyện Bình Gia"
            merged_dict["cap_gcn"]["noi_cap"] = found_auth

        # 6. Fallback: Ngày cấp GCN (GCN_ngayCap)
        curr_ngay = merged_dict.get("cap_gcn", {}).get("ngay_cap", "")
        if not curr_ngay or curr_ngay in ["-", "None"]:
            def fix_digits(cand: str) -> str:
                repl = {'S': '5', 's': '5', 'O': '0', 'o': '0', 'l': '1', 'I': '1', 'i': '1', 'B': '8', 'q': '9', 'A': '1', 'a': '1'}
                for k, v in repl.items():
                    cand = cand.replace(k, v)
                return cand

            for item in result["ocr_lines"]:
                t_line = item.get("text", "")
                if "tiếp nhận" in t_line.lower() or any(k in t_line.lower() for k in ["thời hạn", "mục đích", "đến 11", "đến 12"]):
                    continue
                m_dt = re.search(r'(?:Lộc\s*Bình|Bình\s*Gia|Nam\s*Quan)?[,\s]*(?:ngày|ngay)\s*([0-9A-Za-z]{1,2})\s*(?:tháng|thang)\s*([0-9A-Za-z]{1,2})\s*(?:năm|nam)\s*([0-9/]{2,5})', t_line, re.IGNORECASE)
                if m_dt:
                    d_str = fix_digits(m_dt.group(1))
                    m_str = fix_digits(m_dt.group(2))
                    y_str = fix_digits(m_dt.group(3)).replace('/', '1')
                    if len(y_str) == 2:
                        y_str = '20' + y_str
                    if len(y_str) == 4 and d_str.isdigit() and m_str.isdigit():
                        d_num, mo_num, y_num = int(d_str), int(m_str), int(y_str)
                        if 1 <= d_num <= 31 and 1 <= mo_num <= 12 and 1990 <= y_num <= 2030:
                            merged_dict["cap_gcn"]["ngay_cap"] = f"{d_num:02d}/{mo_num:02d}/{y_num}"
                            break

        # 7. Fallback: Mã vạch GCN (GCN_maVach)
        curr_mv = merged_dict.get("ma_vach", "")
        if not curr_mv or curr_mv in ["-", "None"]:
            for item in reversed(result["ocr_lines"]):
                d_cand = re.sub(r"\D", "", item.get("text", ""))
                if len(d_cand) in [13, 14, 15]:
                    merged_dict["ma_vach"] = d_cand
                    break

        # 8. Fallback: CMND Vợ/Chồng
        chu2_name = merged_dict.get("nguoi_su_dung", {}).get("ho_ten_chu_2", "")
        curr_cid2 = merged_dict.get("nguoi_su_dung", {}).get("cmnd_chu_2", "")
        if chu2_name and (not curr_cid2 or curr_cid2 in ["-", "None"]):
            all_cids = []
            for item in result["ocr_lines"]:
                if item.get("page", 1) <= 3:
                    for m in re.finditer(r"\b(\d{9}|\d{12})\b", item.get("text", "")):
                        c_val = m.group(1)
                        if c_val not in all_cids:
                            all_cids.append(c_val)
            cid1 = merged_dict.get("nguoi_su_dung", {}).get("cmnd_chu_1", "")
            for c_val in all_cids:
                if c_val != cid1:
                    merged_dict["nguoi_su_dung"]["cmnd_chu_2"] = c_val
                    break

        # 9. Fallback: Số thứ tự thửa (TD_soThuTuThua)
        curr_st = merged_dict.get("thua_dat", {}).get("so_thua", "")
        v_st, n_st, _ = GCNValidators.validate_parcel_number(curr_st) if curr_st else (False, "", None)
        if not v_st:
            for item in result["ocr_lines"]:
                t_line = unicodedata.normalize('NFC', item.get("text", ""))
                if any(k in t_line.lower() for k in ["tổng số", "tong so", "nhà ở", "rừng", "cây lâu năm"]):
                    continue
                m = re.search(r'(?:(?:[aâ]\)|ai|[0-9][,\.]?|\-)?\s*)?(?:th[ửừứaảãạuủùú][aàáảãạâầấẩẫậăằắẳẵặc]?\s*(?:đ[ấa]t|[áa]n)?\s*(?:s[oóòỏõọôốồổỗộơớờởỡợ]|bối|sới)|thua\s*(?:dat)?\s*so)\s*[:\.,;\s]\s*(\d+[A-Za-z]?(?:\s*[\+]\s*\d+[A-Za-z]?)*)', t_line, re.IGNORECASE)
                if m:
                    cand = m.group(1).replace(" ", "")
                    v_ok, n_cand, _ = GCNValidators.validate_parcel_number(cand)
                    if v_ok:
                        merged_dict["thua_dat"]["so_thua"] = n_cand
                        break
                m_pre = re.search(r'\b(\d{1,4})[,\s]+(?:tờ\s*(?:bản\s*đồ)?\s*số)', t_line, re.IGNORECASE)
                if m_pre:
                    v_ok, n_cand, _ = GCNValidators.validate_parcel_number(m_pre.group(1))
                    if v_ok:
                        merged_dict["thua_dat"]["so_thua"] = n_cand
                        break

        # 10. Fallback: Số hiệu tờ bản đồ (TD_soHieuToBanDo)
        curr_tb = merged_dict.get("thua_dat", {}).get("to_ban_do", "")
        v_tb, n_tb, _ = GCNValidators.validate_map_sheet(curr_tb) if curr_tb else (False, "", None)
        if not v_tb:
            for item in result["ocr_lines"]:
                t_line = unicodedata.normalize('NFC', item.get("text", ""))
                m = re.search(r'(?:[t|l|v]ờ\s*(?:[bh]ản|[bh]án)?\s*(?:đồ|đổ|đỗ|đó)?\s*(?:số|sới|sốc|so)?|và\s*bán\s*đó\s*(?:số|sới|sốc)?|to\s*ban\s*do\s*so)\s*[:\.,;\s]\s*(\d+(?:\s*[\+]\s*\d+)*)', t_line, re.IGNORECASE)
                if m:
                    v_ok, n_cand, _ = GCNValidators.validate_map_sheet(m.group(1).replace(" ", ""))
                    if v_ok:
                        merged_dict["thua_dat"]["to_ban_do"] = n_cand
                        # Cập nhật cho cả các item trong danh_sach_thua nếu thiếu tờ BĐ
                        for p_item in merged_dict.get("thua_dat", {}).get("danh_sach_thua", []):
                            if not p_item.get("to_ban_do"):
                                p_item["to_ban_do"] = n_cand
                        v_tb = True
                        break

        # 10b. Fallback cặp số trong bảng nhiều thửa (vd: '114' ngay trước số thửa '30')
        if not v_tb and merged_dict.get("thua_dat", {}).get("danh_sach_thua"):
            ds = merged_dict["thua_dat"]["danh_sach_thua"]
            st_set = set(str(p.get("so_thua") or "").strip() for p in ds if p.get("so_thua"))
            ocr_texts = [it.get("text", "").strip() for it in result["ocr_lines"]]
            for idx in range(len(ocr_texts) - 1):
                l1 = ocr_texts[idx]
                l2 = ocr_texts[idx+1]
                if l2 in st_set and l1.isdigit() and len(l1) <= 4:
                    v_ok, n_cand, _ = GCNValidators.validate_map_sheet(l1)
                    if v_ok:
                        merged_dict["thua_dat"]["to_ban_do"] = n_cand
                        for p_item in ds:
                            if not p_item.get("to_ban_do"):
                                p_item["to_ban_do"] = n_cand
                        v_tb = True
                        break

        # 11. Fallback: Diện tích thửa đất (TD_dienTich)
        curr_dt = merged_dict.get("thua_dat", {}).get("dien_tich", "")
        v_dt, n_dt, _ = GCNValidators.validate_area(curr_dt) if curr_dt else (False, None, None)
        if not v_dt:
            for item in result["ocr_lines"]:
                t_line = unicodedata.normalize('NFC', item.get("text", ""))
                if "bằng chữ" in t_line.lower() or "tổng" in t_line.lower():
                    continue
                m_dt = re.search(r'(?:diện\s*tích|dien\s*tich)\s*[:\.]?\s*([\d\.,]+)\s*m', t_line, re.IGNORECASE)
                if m_dt:
                    v_ok, n_cand, _ = GCNValidators.validate_area(m_dt.group(1))
                    if v_ok:
                        merged_dict["thua_dat"]["dien_tich"] = n_cand
                        break

        # 12. Fallback: Thời hạn sử dụng (TD_thoiHanSuDung)
        curr_th = merged_dict.get("thua_dat", {}).get("thoi_han", "")
        v_th, n_th, _ = GCNValidators.validate_land_use_term(curr_th) if curr_th else (False, "", None)
        if not v_th or "mục đích" in str(curr_th).lower():
            found_th = None
            for item in result["ocr_lines"]:
                t_line = unicodedata.normalize('NFC', item.get("text", ""))
                if re.search(r"(?:[eđcCdD][\)\.]?\s*|[-*]\s*)?(?:thời|thành)\s*hạn\s*(?:sử|sắt)?\s*dụng", t_line, re.IGNORECASE):
                    val = re.sub(r"^.*?(?:thời\s*hạn\s*(?:sử\s*dụng)?|thành\s*hạn\s*sắt\s*dụng)\s*[:\.]?\s*", "", t_line, flags=re.IGNORECASE).strip()
                    if val and not any(b in val.lower() for b in ["mục đích", "diện tích", "địa chỉ"]):
                        v_ok, n_cand, _ = GCNValidators.validate_land_use_term(val)
                        if v_ok:
                            found_th = n_cand
                            break
            if not found_th:
                for item in result["ocr_lines"]:
                    t_line = unicodedata.normalize('NFC', item.get("text", ""))
                    if any(k in t_line.lower() for k in ["tiếp nhận", "cmnd", "ngày cấp"]):
                        continue
                    v_ok, n_cand, _ = GCNValidators.validate_land_use_term(t_line)
                    if v_ok:
                        found_th = n_cand
                        break
            if found_th:
                merged_dict["thua_dat"]["thoi_han"] = found_th

        # 13. Fallback: Nguồn gốc sử dụng (TD_nguonGoc)
        curr_ng = merged_dict.get("thua_dat", {}).get("nguon_goc", "")
        v_ng, n_ng, _, _ = GCNValidators.validate_land_use_origin(curr_ng) if curr_ng else (False, "", "", None)
        if not v_ng:
            found_ng = None
            for item in result["ocr_lines"]:
                t_line = unicodedata.normalize('NFC', item.get("text", ""))
                if re.search(r"(?:(?:Thường\s*)?[g8bB][\)\.]?\s*|[-*]\s*)?(?:nguồn\s*gốc\s*(?:sử\s*dụng)?|nguồn\s*gốc:?)", t_line, re.IGNORECASE):
                    val = re.sub(r"^.*?(?:nguồn\s*gốc\s*sử\s*dụng|nguồn\s*gốc)\s*[:\.]?\s*", "", t_line, flags=re.IGNORECASE).strip()
                    if val and not any(b in val.lower() for b in ["người nhận", "tài sản", "thời hạn"]):
                        v_ok, n_cand, _, _ = GCNValidators.validate_land_use_origin(val)
                        if v_ok:
                            found_ng = n_cand
                            break
            if not found_ng:
                for item in result["ocr_lines"]:
                    t_line = unicodedata.normalize('NFC', item.get("text", ""))
                    if any(k in t_line.lower() for k in ["tiếp nhận", "đơn đề nghị"]):
                        continue
                    if any(k in t_line.lower() for k in ["công nhận qsdđ", "giao đất không thu tiền", "giao đất có thu tiền", "qsdd"]):
                        v_ok, n_cand, _, _ = GCNValidators.validate_land_use_origin(t_line)
                        if v_ok:
                            found_ng = n_cand
                            break
            if found_ng:
                merged_dict["thua_dat"]["nguon_goc"] = found_ng

        return result

    @classmethod
    def get_excel_preview_data(cls, record: Dict[str, Any]) -> Dict[str, Any]:
        """
        Tra ve JSON payload xem truoc truc tiep tren Web UI truoc khi tai tep.
        """
        doc_id = record.get("id", "")
        file_name = record.get("file_name", "")
        template = record.get("template", "")
        total_pages = record.get("total_pages", 1)
        created_at = record.get("created_at", "")
        raw_markdown = record.get("raw_markdown", "")

        parsed = cls.parse_raw_markdown(raw_markdown)

        meta = parsed.get("metadata", {})
        meta.setdefault("job_id", doc_id)
        meta.setdefault("file_name", file_name)
        meta.setdefault("template", template)
        meta.setdefault("total_pages", total_pages)
        meta.setdefault("created_at", created_at)

        cadastral_rows = []
        if parsed.get("merged_dict"):
            try:
                rows = Cadastral129Mapper.map_merged_to_rows(
                    parsed["merged_dict"],
                    start_stt=1,
                    file_name=file_name
                )
                if rows:
                    cadastral_rows = rows
            except Exception as e:
                logger.warning(f"Loi sinh preview 129 cot tu raw markdown: {e}")

        return {
            "document_id": doc_id,
            "file_name": file_name,
            "template": template,
            "metadata": meta,
            "summary_fields": parsed.get("summary_fields", []),
            "ocr_lines": parsed.get("ocr_lines", []),
            "cadastral_129_rows": cadastral_rows,
            "total_summary_fields": len(parsed.get("summary_fields", [])),
            "total_ocr_lines": len(parsed.get("ocr_lines", [])),
        }

    @classmethod
    def export_single_record_to_excel(cls, record: Dict[str, Any]) -> bytes:
        """
        Xuat toan bo du lieu tho cua 1 ho so ra tep Excel .xlsx chuyen nghiep (2 Sheet).
        """
        preview = cls.get_excel_preview_data(record)
        meta = preview.get("metadata", {})
        summary_fields = preview.get("summary_fields", [])
        ocr_lines = preview.get("ocr_lines", [])

        wb = openpyxl.Workbook()

        # Sheet 1: Tom Tat Boc Tach Tho
        ws1 = wb.active
        ws1.title = "Tóm Tắt Bóc Tách Thô"
        ws1.views.sheetView[0].showGridLines = True

        font_title = Font(name="Segoe UI", size=14, bold=True, color="1E293B")
        font_meta_k = Font(name="Segoe UI", size=10, bold=True, color="475569")
        font_meta_v = Font(name="Segoe UI", size=10, color="0F172A")
        font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        font_data = Font(name="Segoe UI", size=10, color="1E293B")
        font_data_bold = Font(name="Segoe UI", size=10, bold=True, color="0F172A")

        fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        fill_meta = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        fill_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

        thin_side = Side(border_style="thin", color="CBD5E1")
        border_box = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        # Banner Title
        ws1.merge_cells("A1:C1")
        cell_t = ws1["A1"]
        cell_t.value = "KẾT QUẢ DỮ LIỆU THÔ OCR (RAW OCR DATA)"
        cell_t.font = font_title
        cell_t.alignment = align_left
        ws1.row_dimensions[1].height = 28

        # Metadata Box
        meta_items = [
            ("Mã Hồ Sơ (Job ID)", meta.get("job_id", "-")),
            ("Tên Tệp Hồ Sơ", meta.get("file_name", "-")),
            ("Mẫu Sổ Nhận Diện", meta.get("template", "-")),
            ("Tổng Số Trang", str(meta.get("total_pages", "-"))),
            ("Thời Điểm Quét OCR", meta.get("created_at", "-")),
        ]

        curr_row = 3
        for k, v in meta_items:
            ws1.cell(row=curr_row, column=1, value=k).font = font_meta_k
            ws1.cell(row=curr_row, column=1).fill = fill_meta
            ws1.cell(row=curr_row, column=1).border = border_box
            ws1.cell(row=curr_row, column=2, value=v).font = font_meta_v
            ws1.cell(row=curr_row, column=2).border = border_box
            ws1.cell(row=curr_row, column=3, value="").border = border_box
            ws1.row_dimensions[curr_row].height = 20
            curr_row += 1

        curr_row += 1

        # Table Summary Fields Header
        headers1 = ["STT", "Trường Dữ Liệu Bóc Tách Thô", "Giá Trị Trích Xuất (Raw Value)"]
        for c_idx, h in enumerate(headers1, 1):
            cell = ws1.cell(row=curr_row, column=c_idx, value=h)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_header
            cell.border = border_box
        ws1.row_dimensions[curr_row].height = 24
        curr_row += 1

        # Table Data
        if summary_fields:
            for item in summary_fields:
                is_zebra = (item.get("stt", 0) % 2 == 0)
                row_fill = fill_alt if is_zebra else None

                c1 = ws1.cell(row=curr_row, column=1, value=item.get("stt", 0))
                c1.alignment = align_center
                c1.font = font_data
                c1.border = border_box

                c2 = ws1.cell(row=curr_row, column=2, value=item.get("field", ""))
                c2.alignment = align_left
                c2.font = font_data_bold
                c2.border = border_box

                c3 = ws1.cell(row=curr_row, column=3, value=item.get("value", ""))
                c3.alignment = align_left
                c3.font = font_data
                c3.border = border_box

                if row_fill:
                    c1.fill = row_fill
                    c2.fill = row_fill
                    c3.fill = row_fill

                ws1.row_dimensions[curr_row].height = 22
                curr_row += 1
        else:
            ws1.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=3)
            empty_c = ws1.cell(row=curr_row, column=1, value="Không có dữ liệu bóc tách Section I.")
            empty_c.alignment = align_center
            empty_c.font = font_meta_k
            curr_row += 1

        ws1.column_dimensions["A"].width = 10
        ws1.column_dimensions["B"].width = 34
        ws1.column_dimensions["C"].width = 55

        # Sheet 2: Van Ban OCR Chi Tiet
        ws2 = wb.create_sheet(title="Văn Bản OCR Chi Tiết")
        ws2.views.sheetView[0].showGridLines = True

        ws2.merge_cells("A1:D1")
        cell_t2 = ws2["A1"]
        cell_t2.value = "TOÀN BỘ VĂN BẢN OCR THÔ THEO THỨ TỰ ĐỌC (SPATIAL READING ORDER)"
        cell_t2.font = font_title
        cell_t2.alignment = align_left
        ws2.row_dimensions[1].height = 28

        curr_row2 = 3
        headers2 = ["STT", "Trang", "Tệp Nguồn", "Nội Dung Văn Bản Nhận Dạng Quang Học (OCR)"]
        fill_header2 = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
        for c_idx, h in enumerate(headers2, 1):
            cell = ws2.cell(row=curr_row2, column=c_idx, value=h)
            cell.font = font_header
            cell.fill = fill_header2
            cell.alignment = align_header
            cell.border = border_box
        ws2.row_dimensions[curr_row2].height = 24
        curr_row2 += 1

        if ocr_lines:
            for item in ocr_lines:
                is_zebra = (item.get("stt", 0) % 2 == 0)
                row_fill = fill_alt if is_zebra else None

                c1 = ws2.cell(row=curr_row2, column=1, value=item.get("stt", 0))
                c1.alignment = align_center
                c1.font = font_data
                c1.border = border_box

                c2 = ws2.cell(row=curr_row2, column=2, value=f"Trang {item.get('page', 1)}")
                c2.alignment = align_center
                c2.font = font_data_bold
                c2.border = border_box

                c3 = ws2.cell(row=curr_row2, column=3, value=item.get("file_name", ""))
                c3.alignment = align_left
                c3.font = font_data
                c3.border = border_box

                c4 = ws2.cell(row=curr_row2, column=4, value=item.get("text", ""))
                c4.alignment = align_left
                c4.font = font_data
                c4.border = border_box

                if row_fill:
                    c1.fill = row_fill
                    c2.fill = row_fill
                    c3.fill = row_fill
                    c4.fill = row_fill

                ws2.row_dimensions[curr_row2].height = 20
                curr_row2 += 1
        else:
            ws2.merge_cells(start_row=curr_row2, start_column=1, end_row=curr_row2, end_column=4)
            empty_c = ws2.cell(row=curr_row2, column=1, value="Không có dòng văn bản OCR nào.")
            empty_c.alignment = align_center
            empty_c.font = font_meta_k

        ws2.column_dimensions["A"].width = 10
        ws2.column_dimensions["B"].width = 14
        ws2.column_dimensions["C"].width = 24
        ws2.column_dimensions["D"].width = 65

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    @classmethod
    def export_table_summary_to_excel(cls, records: List[Dict[str, Any]]) -> bytes:
        """
        Xuat danh sach tat ca cac ban ghi du lieu tho sang 1 tep Excel tong hop.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Danh Sách Hồ Sơ Thô"
        ws.views.sheetView[0].showGridLines = True

        font_title = Font(name="Segoe UI", size=14, bold=True, color="1E293B")
        font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        font_data = Font(name="Segoe UI", size=10, color="1E293B")
        font_data_bold = Font(name="Segoe UI", size=10, bold=True, color="0F172A")

        fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        fill_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        thin_side = Side(border_style="thin", color="CBD5E1")
        border_box = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        ws.merge_cells("A1:G1")
        ws["A1"].value = "BẢNG TỔNG HỢP DỮ LIỆU THÔ OCR (RAW OCR RECORDS)"
        ws["A1"].font = font_title
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 28

        curr_row = 3
        headers = [
            "STT", "Mã Hồ Sơ (Job ID)", "Tên Tệp Hồ Sơ", "Thư Mục Kết Quả",
            "Mẫu Nhận Diện", "Số Phát Hành", "Họ Tên Chủ", "Số CMND/CCCD",
            "Số Thửa", "Tờ Bản Đồ", "Diện Tích (m2)", "Thời Gian Lưu"
        ]
        for c_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=curr_row, column=c_idx, value=h)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border_box
        ws.row_dimensions[curr_row].height = 25
        curr_row += 1

        for idx, r in enumerate(records, 1):
            is_zebra = (idx % 2 == 0)
            row_fill = fill_alt if is_zebra else None

            row_values = [
                idx,
                r.get("id", ""),
                r.get("file_name", ""),
                r.get("folder_result", "") or r.get("source_folder", ""),
                r.get("template", ""),
                r.get("so_phat_hanh", ""),
                r.get("ten_chu", ""),
                r.get("cmnd", ""),
                r.get("so_thua", ""),
                r.get("to_ban_do", ""),
                r.get("dien_tich", ""),
                r.get("created_at", "")
            ]

            for c_i, val in enumerate(row_values, 1):
                c = ws.cell(row=curr_row, column=c_i, value=val)
                c.font = font_data_bold if c_i in (3, 7) else font_data
                c.alignment = Alignment(
                    horizontal="left" if c_i in (3, 4, 7) else "center",
                    vertical="center"
                )
                c.border = border_box
                if row_fill:
                    c.fill = row_fill

            ws.row_dimensions[curr_row].height = 22
            curr_row += 1

        widths = [8, 24, 32, 20, 15, 18, 25, 18, 12, 12, 15, 20]
        for idx, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = w

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    @classmethod
    def export_129_from_raw(cls, record: Dict[str, Any]) -> bytes:
        """
        Chuyển đổi các trường bóc tách từ Markdown thô sang bảng Excel 129 cột của Bộ TN&MT.
        """
        import os
        import tempfile

        raw_markdown = record.get("raw_markdown", "")
        file_name = record.get("file_name", "raw_doc.pdf")
        parsed = cls.parse_raw_markdown(raw_markdown)
        merged_dict = parsed.get("merged_dict", {})

        rows = Cadastral129Mapper.map_merged_to_rows(
            merged_dict,
            start_stt=1,
            file_name=file_name
        )
        
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name
        try:
            Excel129Exporter.export(rows, output_path=tmp_path)
            with open(tmp_path, "rb") as f:
                data = f.read()
            return data
        finally:
            if os.path.exists(tmp_path):
                try:
                    os.unlink(tmp_path)
                except Exception:
                    pass
