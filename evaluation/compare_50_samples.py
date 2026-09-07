"""
evaluation/compare_50_samples.py - Chạy kiểm thử 50 mẫu ngẫu nhiên từ ClearData và đối soát 3 bên:
1. Bên 1: Metadata / Ground Truth từ Cây thư mục (Tờ, Thửa, Tên chủ hồ sơ).
2. Bên 2: Kết quả OCR Pipeline (2D Spatial Engine + Modular Domain Parsers) ở chế độ OCR_ONLY (Zero-Leakage).
3. Bên 3: Bảng đối soát chéo 3 bên (Exact match số thửa, Entity-level match tên chủ).

Hỗ trợ chạy song song đa tiến trình (Multiprocessing) để hoàn thành toàn bộ 50 mẫu trong vài phút.
"""

import gc
import json
import os
import random
import re
import sys
import time
import multiprocessing as mp
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Đảm bảo in UTF-8 không lỗi trên Windows
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

# Chuyển toàn bộ Cache sang ổ D
os.environ.setdefault("HF_HOME", r"D:\Tho\OCR\.cache\huggingface")
os.environ.setdefault("HUGGINGFACE_HUB_CACHE", r"D:\Tho\OCR\.cache\huggingface\hub")
os.environ.setdefault("TRANSFORMERS_CACHE", r"D:\Tho\OCR\.cache\huggingface\transformers")
os.environ.setdefault("TORCH_HOME", r"D:\Tho\OCR\.cache\torch")
os.environ.setdefault("PADDLE_HOME", r"D:\Tho\OCR\.cache\paddle")

PROJECT_ROOT = Path(__file__).parent.parent.resolve()
sys.path.insert(0, str(PROJECT_ROOT))

import logging
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s"
)
logging.getLogger("ppocr").setLevel(logging.WARNING)

from preprocessing.ingestion import Ingestion
from preprocessing.deskew import Deskew
from preprocessing.color_profile import ColorProfile
from preprocessing.seal_mask import SealMask
from preprocessing.orientation import OrientationCorrector
from detection.paddleocr_detect import PaddleOCRDetector
from recognition.vietocr_recognize import VietOCRRecognizer
from extraction.template_classifier import TemplateClassifier
from extraction.page_grouper import PageGrouper
from extraction.label_anchor_extractor import LabelAnchorExtractor
from extraction.cross_validate import CrossValidate
from extraction.diagram_extractor import DiagramExtractor
from extraction.address_normalizer import AddressNormalizer
from extraction.gcn_merger import GCNMerger
from evaluation.ground_truth_evaluator import GroundTruthEvaluator
from api.main import run_pipeline_on_image
from process_cleardata_batch import parse_folder_metadata, sanitize_filename

# Global worker pipeline for multiprocessing
worker_pipeline = None


def init_worker():
    global worker_pipeline
    try:
        import torch
        torch.set_num_threads(2)
    except Exception:
        pass
    worker_pipeline = {
        "ingestion": Ingestion(),
        "deskew": Deskew(),
        "color_profile": ColorProfile(str(PROJECT_ROOT / "configs" / "color_profiles.json")),
        "seal_mask": SealMask(str(PROJECT_ROOT / "configs" / "color_profiles.json")),
        "detector": PaddleOCRDetector(use_gpu=False),
        "recognizer": VietOCRRecognizer(device="cpu"),
        "classifier": TemplateClassifier(),
        "page_grouper": PageGrouper(str(PROJECT_ROOT / "configs" / "template_labels.json")),
        "extractor": LabelAnchorExtractor(str(PROJECT_ROOT / "configs" / "template_labels.json")),
        "cross_validator": CrossValidate(),
        "diagram_extractor": DiagramExtractor(),
        "address_normalizer": AddressNormalizer(),
    }
    gc.collect()


def evaluate_3_way(ocr_res: Dict[str, Any], meta: Dict[str, Any]) -> Dict[str, Any]:
    """
    So sánh đối soát 3 bên giữa Ground Truth thư mục và Kết quả OCR Pipeline
    sử dụng GroundTruthEvaluator (chuẩn hóa exact số thửa và entity-level tên chủ).
    """
    eval_res = GroundTruthEvaluator.evaluate_sample(ocr_res, meta)
    thua_dict = ocr_res.get("thua_dat", {})
    chu_dict = ocr_res.get("nguoi_su_dung", {})
    cap_dict = ocr_res.get("cap_gcn", {})

    return {
        "to_ban_do": eval_res["to_ban_do"],
        "so_thua": eval_res["so_thua"],
        "ten_chu": {
            "gt": eval_res["ten_chu"]["gt"],
            "ocr": eval_res["ten_chu"]["ocr"],
            "match": eval_res["ten_chu"]["match"],
            "score": eval_res["ten_chu"]["similarity_score"]
        },
        "so_phat_hanh": ocr_res.get("so_phat_hanh", ""),
        "so_vao_so": ocr_res.get("so_vao_so", ""),
        "dien_tich": thua_dict.get("dien_tich_cap", ""),
        "cccd": chu_dict.get("cmnd_chu_1", "") or chu_dict.get("cmnd", ""),
        "ngay_sinh": chu_dict.get("ngay_sinh_chu_1", "") or chu_dict.get("ngay_sinh", ""),
        "noi_cap": cap_dict.get("noi_cap", ""),
        "ngay_cap": cap_dict.get("ngay_cap", ""),
        "nguoi_ky": cap_dict.get("nguoi_ky_qd", ""),
        "muc_dich": thua_dict.get("muc_dich_su_dung", ""),
        "ma_muc_dich": thua_dict.get("ma_muc_dich", ""),
        "is_all_core_match": eval_res["is_all_core_match"]
    }


def process_file_task(task_args: Tuple[int, str, str]) -> Tuple[int, Dict[str, Any], Optional[Dict[str, Any]], Optional[str]]:
    """Xử lý 1 file duy nhất cho 1 worker process."""
    idx, f_path_str, data_dir_str = task_args
    global worker_pipeline
    f_path = Path(f_path_str)
    data_dir = Path(data_dir_str)
    rel_path = f_path.relative_to(data_dir)
    meta = parse_folder_metadata(f_path, data_dir)
    owner_clean = sanitize_filename(meta.get("ten_chu_thu_muc", "")) or "Unknown"
    to_str = f"To_{meta.get('to_ban_do')}" if meta.get('to_ban_do') else "To_X"
    thua_str = f"Thua_{meta.get('so_thua')}" if meta.get('so_thua') else "Thua_Y"
    doc_id = f"{to_str}_{thua_str}_{owner_clean}_{idx}"

    t_f = time.time()
    try:
        pages = worker_pipeline["ingestion"].load(str(f_path), split_a3=True, smart_gcn_filter=True)
        if not pages:
            pages = worker_pipeline["ingestion"].load(str(f_path), split_a3=False, smart_gcn_filter=False)

        page_results = []
        for p_idx, page_img in enumerate(pages, 1):
            page_job_id = f"{doc_id}_p{p_idx}"
            p_res = run_pipeline_on_image(page_img, worker_pipeline, page_job_id)
            p_res["file_name"] = f"{f_path.stem}_p{p_idx}.png"
            page_results.append(p_res)

        # Chế độ ocr_only=True: KHÔNG DÙNG metadata thư mục (Zero-Leakage)
        merged = GCNMerger.merge(page_results, bo_gcn_id=doc_id, folder_meta=None, ocr_only=True)
        elapsed = round(time.time() - t_f, 2)
        merged["tong_thoi_gian_sec"] = elapsed
        merged["source_file"] = str(rel_path)
        merged["folder_meta"] = meta

        # Đối soát 3 bên với Ground Truth
        cmp_res = evaluate_3_way(merged, meta)
        merged["doi_soat_3_ben"] = cmp_res

        del pages, page_results
        gc.collect()
        return idx, merged, cmp_res, None
    except Exception as exc:
        elapsed = round(time.time() - t_f, 2)
        err_res = {
            "bo_gcn": doc_id,
            "source_file": str(rel_path),
            "folder_meta": meta,
            "error": str(exc),
            "tong_thoi_gian_sec": elapsed
        }
        gc.collect()
        return idx, err_res, None, str(exc)


def main():
    mp.freeze_support()

    print("=" * 110)
    print("      ĐỐI SOÁT 3 BÊN: 50 MẪU NGẪU NHIÊN CLEARDATA (PIPELINE OCR 2D SPATIAL ENGINE)")
    print("                      CHẾ ĐỘ OCR_ONLY = TRUE (ZERO METADATA LEAKAGE)            ")
    print("=" * 110)

    data_dir = Path(r"D:\Tho\OCR\DataOCR\Du Hang sau VILG\ClearData")
    if not data_dir.exists():
        print(f"[LỖI] Không tìm thấy thư mục: {data_dir}")
        return

    output_dir = PROJECT_ROOT / "output"
    output_dir.mkdir(parents=True, exist_ok=True)

    # 1. Thu thập danh sách file
    all_files = sorted(
        list(data_dir.rglob("*.pdf")) +
        list(data_dir.rglob("*.png")) +
        list(data_dir.rglob("*.jpg")) +
        list(data_dir.rglob("*.jpeg"))
    )

    import argparse
    parser = argparse.ArgumentParser(description="Chạy đối soát 3 bên")
    parser.add_argument("--samples", type=int, default=50, help="Số lượng mẫu kiểm thử (mặc định: 50)")
    parser.add_argument("--indices", type=str, default="", help="Chỉ định danh sách STT mẫu cần chạy (ví dụ: 1,2,3,7,10,13)")
    parser.add_argument("--workers", type=int, default=2, help="Số tiến trình worker chạy song song (mặc định: 2 để tránh OOM)")
    args = parser.parse_args()

    SAMPLE_SIZE = min(50, len(all_files))
    random.seed(42)
    full_50_files = sorted(random.sample(all_files, SAMPLE_SIZE), key=lambda x: str(x))

    if args.indices:
        target_idxs = [int(x.strip()) for x in args.indices.split(",") if x.strip()]
        selected_files = [(i, full_50_files[i - 1]) for i in target_idxs if 1 <= i <= len(full_50_files)]
    else:
        limit = min(args.samples, len(full_50_files))
        selected_files = [(i + 1, full_50_files[i]) for i in range(limit)]

    print(f"\n[+] Tổng số file có sẵn trong ClearData: {len(all_files)}")
    print(f"[+] Số lượng hồ sơ thực hiện kiểm thử : {len(selected_files)} / 50")
    print(f"[+] Số worker processes chạy song song  : {args.workers}")

    tasks = [(idx, str(f_path), str(data_dir)) for idx, f_path in selected_files]

    t_start = time.time()
    results_dict: Dict[int, Dict[str, Any]] = {}
    comparisons_dict: Dict[int, Dict[str, Any]] = {}
    success_count = 0

    if args.workers <= 1:
        # Chạy tuần tự (single-process)
        init_worker()
        for t in tasks:
            idx, merged, cmp_res, err = process_file_task(t)
            if err is None:
                success_count += 1
                results_dict[idx] = merged
                comparisons_dict[idx] = cmp_res
                m_to = "✓" if cmp_res["to_ban_do"]["match"] else "✗"
                m_thua = "✓" if cmp_res["so_thua"]["match"] else "✗"
                m_chu = "✓" if cmp_res["ten_chu"]["match"] else "✗"
                chu_val = merged.get("nguoi_su_dung", {}).get("ten", "")
                thua_val = merged.get("thua_dat", {}).get("so_thua", "")
                to_val = merged.get("thua_dat", {}).get("to_ban_do", "")
                sph_val = merged.get("so_phat_hanh", "")
                dt_val = merged.get("thua_dat", {}).get("dien_tich_cap", "")
                elapsed = merged.get("tong_thoi_gian_sec", 0)
                print(f"[{idx:02d}/50] ({elapsed:4.1f}s) | Tờ:{to_val:<3} [{m_to}] | Thửa:{thua_val:<4} [{m_thua}] | Chủ: {chu_val[:25]:<25} [{m_chu}] | Serial: {sph_val:<10} | DT: {dt_val} m²")
            else:
                results_dict[idx] = merged
                print(f"[{idx:02d}/50] | ✗ LỖI: {err}")
    else:
        # Chạy song song đa tiến trình
        print(f"\n[+] Đang khởi tạo pool {args.workers} workers và nạp mô hình OCR...")
        with mp.Pool(processes=args.workers, initializer=init_worker) as pool:
            for idx, merged, cmp_res, err in pool.imap_unordered(process_file_task, tasks):
                if err is None:
                    success_count += 1
                    results_dict[idx] = merged
                    comparisons_dict[idx] = cmp_res
                    m_to = "✓" if cmp_res["to_ban_do"]["match"] else "✗"
                    m_thua = "✓" if cmp_res["so_thua"]["match"] else "✗"
                    m_chu = "✓" if cmp_res["ten_chu"]["match"] else "✗"
                    chu_val = merged.get("nguoi_su_dung", {}).get("ten", "")
                    thua_val = merged.get("thua_dat", {}).get("so_thua", "")
                    to_val = merged.get("thua_dat", {}).get("to_ban_do", "")
                    sph_val = merged.get("so_phat_hanh", "")
                    dt_val = merged.get("thua_dat", {}).get("dien_tich_cap", "")
                    elapsed = merged.get("tong_thoi_gian_sec", 0)
                    print(f"[{idx:02d}/50] ({elapsed:4.1f}s) | Tờ:{to_val:<3} [{m_to}] | Thửa:{thua_val:<4} [{m_thua}] | Chủ: {chu_val[:25]:<25} [{m_chu}] | Serial: {sph_val:<10} | DT: {dt_val} m²")
                else:
                    results_dict[idx] = merged
                    print(f"[{idx:02d}/50] | ✗ LỖI: {err}")

    total_time = round(time.time() - t_start, 2)

    # Sắp xếp lại danh sách kết quả theo đúng thứ tự STT 1..N
    sorted_indices = sorted(results_dict.keys())
    results = [results_dict[i] for i in sorted_indices]
    comparisons = [comparisons_dict[i] for i in sorted_indices if i in comparisons_dict]

    # 4. Thống kê tổng hợp
    print("\n" + "=" * 110)
    print("                       TỔNG KẾT ĐỐI SOÁT 3 BÊN TRÊN CLEARDATA")
    print("=" * 110)
    total_valid = len(comparisons)
    to_matches = sum(1 for c in comparisons if c["to_ban_do"]["match"])
    thua_matches = sum(1 for c in comparisons if c["so_thua"]["match"])
    chu_matches = sum(1 for c in comparisons if c["ten_chu"]["match"])
    all_core_matches = sum(1 for c in comparisons if c["is_all_core_match"])

    has_serial = sum(1 for c in comparisons if c["so_phat_hanh"])
    has_area = sum(1 for c in comparisons if c["dien_tich"])
    has_cccd = sum(1 for c in comparisons if c["cccd"])
    has_issue_place = sum(1 for c in comparisons if c["noi_cap"])
    has_signee = sum(1 for c in comparisons if c["nguoi_ky"])

    n_sel = len(selected_files)
    print(f"• Tổng số hồ sơ kiểm thử            : {n_sel}")
    print(f"• Xử lý thành công                  : {success_count}/{n_sel} ({success_count/n_sel*100:.1f}%)" if n_sel > 0 else "")
    print(f"• Tổng thời gian thực thi            : {total_time:.1f}s ({total_time/60:.2f} phút)")
    print(f"• Thời gian trung bình / hồ sơ      : {total_time/n_sel:.2f}s\n" if n_sel > 0 else "")

    print("-" * 80)
    print(f"  {'HẠNG MỤC ĐỐI SOÁT 3 BÊN':<35} | {'SỐ LƯỢNG ĐẠT':<15} | {'TỶ LỆ KHỚP'}")
    print("-" * 80)
    metrics = [
        ("1. Khớp Tờ bản đồ (Bên 1 ↔ Bên 2)", to_matches, total_valid),
        ("2. Khớp Số thửa đất (Bên 1 ↔ Bên 2)", thua_matches, total_valid),
        ("3. Khớp Tên chủ sử dụng (Bên 1 ↔ Bên 2)", chu_matches, total_valid),
        ("4. Khớp TOÀN DIỆN cả Tờ + Thửa + Chủ", all_core_matches, total_valid),
        ("5. Bóc tách được Số phát hành (Serial)", has_serial, total_valid),
        ("6. Bóc tách được Diện tích cấp (m2)", has_area, total_valid),
        ("7. Bóc tách được CCCD / CMND        ", has_cccd, total_valid),
        ("8. Bóc tách được Nơi cấp GCN        ", has_issue_place, total_valid),
        ("9. Bóc tách được Người ký quyết định", has_signee, total_valid),
    ]
    for label, cnt, tot in metrics:
        pct = cnt / tot * 100.0 if tot > 0 else 0.0
        bar = "█" * int(pct / 5) + "░" * (20 - int(pct / 5))
        print(f"  {label:<35} | {cnt:02d}/{tot:02d} [{bar}] | {pct:5.1f}%")

    # 5. Xuất file Excel đối soát 3 bên
    excel_path = output_dir / "SO_SANH_3_BEN_50_MAU.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Doi_Soat_3_Ben_50_Mau"

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_fill_eval = PatternFill(start_color="276A3C", end_color="276A3C", fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    border_thin = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                         top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))
    pass_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fail_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    headers = [
        "STT", "Hồ sơ Scan / PDF",
        "GT: Tờ BĐ", "OCR: Tờ BĐ", "Khớp Tờ",
        "GT: Số Thửa", "OCR: Số Thửa", "Khớp Thửa",
        "GT: Tên Chủ Thư Mục", "OCR: Chủ Sử Dụng", "Khớp Chủ", "Độ Tương Đồng (%)",
        "1. Mẫu Sổ", "2. Số Phát Hành (Serial)", "3. Số Vào Sổ",
        "4. Tỷ Lệ Bản Đồ", "5. Diện Tích Cấp (m2)", "6. Diện Tích Bằng Chữ",
        "7. Mục Đích Sử Dụng", "8. Mã Mục Đích",
        "9. CCCD/CMND", "10. Năm Sinh", "11. Nơi Cấp GCN",
        "12. Ngày Cấp GCN", "13. Người Ký", "Thời Gian Xử Lý (s)"
    ]
    ws.append(headers)
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(1, col_idx)
        cell.font = header_font
        cell.alignment = align_center
        cell.fill = header_fill_eval if ("GT:" in h or "Khớp" in h) else header_fill

    for idx_row, r in enumerate(results, 1):
        meta = r.get("folder_meta", {})
        cmp = r.get("doi_soat_3_ben", {})
        thua = r.get("thua_dat", {})
        chu = r.get("nguoi_su_dung", {})
        cap = r.get("cap_gcn", {})

        to_match = "ĐẠT" if cmp.get("to_ban_do", {}).get("match") else "LỆCH"
        thua_match = "ĐẠT" if cmp.get("so_thua", {}).get("match") else "LỆCH"
        chu_match = "ĐẠT" if cmp.get("ten_chu", {}).get("match") else "LỆCH"
        score_chu = f"{cmp.get('ten_chu', {}).get('score', 0):.1f}%"

        row_data = [
            idx_row,
            r.get("source_file", ""),
            meta.get("to_ban_do", ""),
            thua.get("to_ban_do", ""),
            to_match,
            meta.get("so_thua", ""),
            thua.get("so_thua", ""),
            thua_match,
            meta.get("ten_chu_thu_muc", ""),
            chu.get("ten", ""),
            chu_match,
            score_chu,
            r.get("mau", ""),
            r.get("so_phat_hanh", ""),
            r.get("so_vao_so", ""),
            thua.get("ty_le", ""),
            thua.get("dien_tich_cap", ""),
            thua.get("dien_tich_chu", ""),
            thua.get("muc_dich_su_dung", ""),
            thua.get("ma_muc_dich", ""),
            chu.get("cmnd", ""),
            chu.get("ngay_sinh", ""),
            cap.get("noi_cap", ""),
            cap.get("ngay_cap", ""),
            cap.get("nguoi_ky_qd", ""),
            r.get("tong_thoi_gian_sec", 0)
        ]
        ws.append(row_data)
        curr_row = ws.max_row

        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(curr_row, col_idx)
            c.border = border_thin
            if col_idx in [5, 8, 11]:
                c.alignment = align_center
                c.fill = pass_fill if c.value == "ĐẠT" else fail_fill

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(str(excel_path))
    print(f"\n[✓] Đã xuất file Excel đối soát 3 bên tại: {excel_path}")

    # Xuất JSON
    json_path = output_dir / "so_sanh_3_ben_50_mau.json"
    json_path.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"[✓] Đã lưu file JSON chi tiết tại: {json_path}")
    print("=" * 110)

    # Tự động xuất bảng Full 29 trường chuẩn hóa
    try:
        import subprocess
        subprocess.run([sys.executable, str(PROJECT_ROOT / "evaluation" / "export_full_29_fields_excel.py")], check=True)
    except Exception as exc:
        print(f"Cảnh báo khi xuất bảng 29 trường: {exc}")


if __name__ == "__main__":
    main()
