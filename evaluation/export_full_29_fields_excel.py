"""
evaluation/export_full_29_fields_excel.py - Xuất bảng Excel 29 trường chuẩn hóa GCNSchemaV1.

Cố định đúng 29 trường địa chính (không đánh số tới 34), bổ sung cột Hàng đợi Review.
"""

import json
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).parent.parent.resolve()
output_dir = PROJECT_ROOT / "output"
json_path = output_dir / "so_sanh_3_ben_50_mau.json"

if not json_path.exists():
    print(f"Không tìm thấy file: {json_path}")
    exit(1)

with open(json_path, "r", encoding="utf-8") as f:
    results = json.load(f)

excel_path = output_dir / "SO_SANH_3_BEN_50_MAU_FULL_29_TRUONG.xlsx"
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Doi_Soat_29_Truong_50_Mau"

# Styling definitions
header_fill_ocr = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_fill_eval = PatternFill(start_color="276A3C", end_color="276A3C", fill_type="solid")
header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
border_thin = Border(
    left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
)
pass_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
fail_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")

# Đúng 29 cột nghiệp vụ chuẩn
headers = [
    # 0. STT & File
    "STT", "Hồ sơ Scan / PDF",
    
    # 1. Đối soát 3 bên cốt lõi
    "GT: Tờ BĐ", "OCR: Tờ BĐ", "Khớp Tờ",
    "GT: Số Thửa", "OCR: Số Thửa", "Khớp Thửa",
    "GT: Tên Chủ Thư Mục", "OCR: Chủ Sử Dụng", "Khớp Chủ", "Độ Tương Đồng (%)",
    
    # 2. Nhóm Định danh GCN (3 trường)
    "1. Số Phát Hành (Serial Phôi)", "2. Số Vào Sổ Cấp GCN", "3. Mã Vạch",
    
    # 3. Nhóm Chủ Sử Dụng & Nhân Thân (8 trường)
    "4. Chủ 1 (Chồng/Đại diện)", "5. CCCD/CMND Chủ 1", "6. Năm Sinh Chủ 1",
    "7. Chủ 2 (Vợ/Đồng SH)", "8. CCCD/CMND Chủ 2", "9. Năm Sinh Chủ 2",
    "10. Địa Chỉ Thường Trú", "11. Loại Chủ Sử Dụng",
    
    # 4. Nhóm Thửa Đất (10 trường)
    "12. Số Thửa Đất", "13. Số Tờ Bản Đồ", "14. Địa Chỉ Thửa Đất",
    "15. Diện Tích Cấp (m2)", "16. Diện Tích Riêng", "17. Diện Tích Chung",
    "18. Diện Tích Bằng Chữ", "19. Mục Đích Sử Dụng", "20. Mã Mục Đích (ODT/ONT)",
    "21. Thời Hạn Sử Dụng",
    
    # 5. Nhóm Cấp Giấy Chứng Nhận (7 trường)
    "22. Hình Thức Sử Dụng", "23. Nguồn Gốc Sử Dụng", "24. Nơi Cấp (Cơ Quan)",
    "25. Ngày Cấp GCN", "26. Người Ký Quyết Định", "27. Chức Vụ Người Ký",
    "28. Tỷ Lệ Bản Đồ",
    
    # 6. Nhóm Biến Động (1 trường)
    "29. Biến Động Mới Nhất",
    
    # 7. Quality & Metrics
    "Hàng Đợi Review (Cảnh Báo)", "Thời Gian Xử Lý (s)"
]

ws.append(headers)
for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(1, col_idx)
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if "GT:" in h or "Khớp" in h:
        cell.fill = header_fill_eval
    else:
        cell.fill = header_fill_ocr

# Điền dữ liệu 50 mẫu
for idx, r in enumerate(results, 1):
    meta = r.get("folder_meta", {})
    cmp = r.get("doi_soat_3_ben", {})
    thua = r.get("thua_dat", {})
    chu = r.get("nguoi_su_dung", {})
    cap = r.get("cap_gcn", {})
    bd = r.get("bien_dong", {})

    to_match = "ĐẠT" if cmp.get("to_ban_do", {}).get("match") else "LỆCH"
    thua_match = "ĐẠT" if cmp.get("so_thua", {}).get("match") else "LỆCH"
    chu_match = "ĐẠT" if cmp.get("ten_chu", {}).get("match") else "LỆCH"
    score_chu = f"{cmp.get('ten_chu', {}).get('score', 0):.1f}%"

    row_data = [
        idx,
        r.get("source_file", ""),
        
        # Đối soát 3 bên
        meta.get("to_ban_do", ""),
        thua.get("to_ban_do", ""),
        to_match,
        meta.get("so_thua", ""),
        thua.get("so_thua", ""),
        thua_match,
        meta.get("ten_chu_thu_muc", ""),
        chu.get("ten", ""),
        chu_match,
        score_chu,
        
        # 1-3. Định danh (3 trường)
        r.get("so_phat_hanh", ""),
        r.get("so_vao_so", ""),
        r.get("ma_vach", ""),
        
        # 4-11. Nhân thân (8 trường)
        chu.get("ho_ten_chu_1", ""),
        chu.get("cmnd_chu_1", ""),
        chu.get("ngay_sinh_chu_1", ""),
        chu.get("ho_ten_chu_2", "") or "N/A",
        chu.get("cmnd_chu_2", "") or "N/A",
        chu.get("ngay_sinh_chu_2", "") or "N/A",
        chu.get("dia_chi_thuong_tru", ""),
        chu.get("loai_chu", "Cá nhân"),
        
        # 12-21. Thửa đất (10 trường)
        thua.get("so_thua", ""),
        thua.get("to_ban_do", ""),
        thua.get("dia_chi", ""),
        thua.get("dien_tich_cap", ""),
        thua.get("dien_tich_rieng", ""),
        thua.get("dien_tich_chung", ""),
        thua.get("dien_tich_chu", ""),
        thua.get("muc_dich_su_dung", ""),
        thua.get("ma_muc_dich", ""),
        thua.get("thoi_han", ""),
        
        # 22-28. Cấp giấy (7 trường)
        thua.get("hinh_thuc_su_dung", ""),
        thua.get("nguon_goc", ""),
        cap.get("noi_cap", ""),
        cap.get("ngay_cap", ""),
        cap.get("nguoi_ky_qd", ""),
        cap.get("chuc_vu_nguoi_ky", ""),
        thua.get("ty_le", ""),
        
        # 29. Biến động (1 trường)
        bd.get("thong_tin_bien_dong", "") or r.get("ten_chuyen_nhuong_moi", "") or "N/A",
        
        # Hàng đợi review
        ", ".join(r.get("can_review", [])),
        r.get("tong_thoi_gian_sec", 0)
    ]
    ws.append(row_data)
    curr_row = ws.max_row

    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(curr_row, col_idx)
        c.border = border_thin
        c.font = Font(name="Segoe UI", size=9)
        if col_idx in [5, 8, 11]:  # Các cột ĐỐI SOÁT ĐẠT/LỆCH
            c.alignment = align_center
            c.fill = pass_fill if c.value == "ĐẠT" else fail_fill

# Tự động căn chỉnh độ rộng cột
for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(min(max_len + 3, 35), 11)

ws.freeze_panes = "C2"
wb.save(str(excel_path))
print(f"Đã xuất thành công bảng Excel 29 trường chuẩn hóa: {excel_path}")
