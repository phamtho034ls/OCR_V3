import os
import sys
import tempfile
import json
from pathlib import Path

# Add project root to sys.path
PROJECT_ROOT = Path(r"d:\Tho\OCR\OCR_V3\ocr-so-do").resolve()
sys.path.insert(0, str(PROJECT_ROOT))

import openpyxl
from extraction.excel_chuyen_doi_mapper import ExcelChuyenDoiMapper
from extraction.excel_template_exporter import ExcelTemplateExporter


def test_columns_json():
    col_path = next(p for p in [
        PROJECT_ROOT / "backend" / "configs" / "excel_chuyen_doi_columns.json",
        PROJECT_ROOT / "configs" / "excel_chuyen_doi_columns.json",
    ] if p.exists())
    with open(col_path, "r", encoding="utf-8") as f:
        cols = json.load(f)
    assert len(cols) == 129, f"Cần đúng 129 cột, thực tế có {len(cols)}"
    assert cols[0]["code"] == "STT"
    assert cols[10]["code"] == "GCN_soPhatHanh"
    for c in cols:
        assert "col" in c and "code" in c and "section" in c and "section_key" in c


def test_mapper_and_exporter():
    mock_merged = {
        "bo_gcn": "TEST_HOSO_001",
        "so_phat_hanh": "CH 123456",
        "so_vao_so": "CS 00987",
        "ngay_cap": "15/08/2021",
        "thua_dat": {
            "so_thua": "125",
            "to_ban_do": "45",
            "dien_tich_cap": "180.5",
            "dien_tich_rieng": "180.5",
            "dien_tich_chung": "0",
            "muc_dich_sd": "Đất ở tại đô thị (ODT)",
            "nguon_goc_sd": "Nhà nước giao đất có thu tiền sử dụng đất",
            "dia_chi": "Số 12, Phố Bà Triệu, Phường Liên Bảo, Thành phố Vĩnh Yên, Tỉnh Vĩnh Phúc"
        },
        "nguoi_su_dung": {
            "ten": "Ông: Nguyễn Văn An, vợ là Bà: Trần Thị Bình",
            "cmnd": "026090001234, 026192005678",
            "ngay_sinh": "1980, 1982",
            "dia_chi": "Số 12, Phố Bà Triệu, Phường Liên Bảo, TP Vĩnh Yên, Tỉnh Vĩnh Phúc"
        }
    }

    row = ExcelChuyenDoiMapper.map_merged_to_row(mock_merged, stt=1, file_name="TEST_HOSO_001.pdf")
    
    assert row["STT"] == 1
    assert row["GCN_soPhatHanh"] == "CH 123456"
    assert row["GCN_soVaoSo"] == "CS 00987"
    assert row["TD_soThuTuThua"] == "125"
    assert row["TD_soHieuToBanDo"] == "45"
    assert row["TD_dienTich"] == 180.5
    assert row["TD_maMucDichSuDung"] == "ODT"
    assert row["CHU_tenXa"] == "Liên Bảo"
    assert row["CHU_tenHuyen"] == "Vĩnh Yên"
    assert row["CHU_tenTinh"] == "Vĩnh Phúc"
    
    assert "Nguyễn Văn An" in row["CHU_hoTen"]
    assert row["CHU_gioiTinh"] == 1
    assert row["GT_soGiayTo"] == "026090001234"
    assert row["GT_loaiGiayTo"] == "CCCD"

    assert "Trần Thị Bình" in row["VC_hoTen"]
    assert row["VC_gioiTinh"] == 0
    assert row["GT_VC_soGiayTo"] == "026192005678"

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        out_path = tmp.name

    try:
        res_file = ExcelTemplateExporter.export([row], out_path)
        assert os.path.exists(res_file)
        
        wb = openpyxl.load_workbook(res_file)
        sheet = wb["KeKhaiDangKy"] if "KeKhaiDangKy" in wb.sheetnames else wb.active
        assert sheet.max_row == 5, f"Sau khi xuất 1 dòng, max_row phải là 5, thực tế là {sheet.max_row}"
        assert sheet.cell(5, 1).value == 1
        assert sheet.cell(5, 11).value == "CH 123456"
        wb.close()
    finally:
        if os.path.exists(out_path):
            os.unlink(out_path)


def test_map_merged_to_rows_multi_parcel():
    mock_multi = {
        "bo_gcn": "BH_405659",
        "so_phat_hanh": "BH 405659",
        "so_vao_so": "CH0003",
        "ngay_cap": "01/12/2011",
        "thua_dat": {
            "so_thua": "4+5+6+8+10+11",
            "to_ban_do": "91",
            "ma_muc_dich": "HNK",
            "muc_dich_su_dung": "Đất bằng trồng cây hàng năm khác",
            "thoi_han": "Đến 12/2031",
            "nguon_goc": "Công nhận QSDĐ như giao đất không thu tiền sử dụng đất",
            "danh_sach_thua": [
                {"so_thua": "4", "to_ban_do": "91", "dien_tich": 243.0, "ma_muc_dich": "HNK", "thoi_han": "Đến 12/2031"},
                {"so_thua": "5", "to_ban_do": "91", "dien_tich": 99.0, "ma_muc_dich": "HNK", "thoi_han": "Đến 12/2031"},
                {"so_thua": "6", "to_ban_do": "91", "dien_tich": 104.0, "ma_muc_dich": "HNK", "thoi_han": "Đến 12/2031"},
                {"so_thua": "8", "to_ban_do": "91", "dien_tich": 110.0, "ma_muc_dich": "HNK", "thoi_han": "Đến 12/2031"},
                {"so_thua": "10", "to_ban_do": "91", "dien_tich": 259.0, "ma_muc_dich": "HNK", "thoi_han": "Đến 12/2031"},
                {"so_thua": "11", "to_ban_do": "91", "dien_tich": 242.0, "ma_muc_dich": "HNK", "thoi_han": "Đến 12/2031"},
            ]
        },
        "nguoi_su_dung": {
            "ten": "Triệu Văn Khé",
            "cmnd": "111000222",
            "ngay_sinh": "1960",
            "dia_chi": "Thôn Nà Làng, xã Vĩnh Yên, huyện Bình Gia, tỉnh Lạng Sơn"
        }
    }

    rows = ExcelChuyenDoiMapper.map_merged_to_rows(mock_multi, start_stt=7, file_name="BH_405659")
    assert len(rows) == 6, f"Phải tách thành đúng 6 dòng, thực tế có {len(rows)} dòng"

    expected_parcels = [("4", 243.0), ("5", 99.0), ("6", 104.0), ("8", 110.0), ("10", 259.0), ("11", 242.0)]
    for i, r in enumerate(rows):
        exp_st, exp_dt = expected_parcels[i]
        assert r["STT"] == 7 + i
        assert r["DDK_maDon"] == f"DON_{7 + i}"
        assert r["GCN_soPhatHanh"] == "BH 405659"
        assert r["GCN_soVaoSo"] == "CH0003"
        assert r["CHU_hoTen"] == "Triệu Văn Khé"
        assert r["TD_soThuTuThua"] == exp_st
        assert r["TD_soHieuToBanDo"] == "91"
        assert r["TD_dienTich"] == exp_dt
        assert r["TD_maMucDichSuDung"] == "HNK"
        assert r["TD_thoiHanSuDung"] == "Đến 12/2031"
        assert "+" not in str(r["TD_soThuTuThua"])


def test_api_endpoints():
    from fastapi.testclient import TestClient
    from api.main import app

    client = TestClient(app)

    # 1. Test GET /chuyen-doi/columns
    res_cols = client.get("/chuyen-doi/columns")
    assert res_cols.status_code == 200, res_cols.text
    cols_data = res_cols.json()
    assert cols_data["total"] == 129
    assert len(cols_data["sections"]) >= 10

    # 2. Test GET /chuyen-doi/vinhyen-50
    res_vy = client.get("/chuyen-doi/vinhyen-50")
    if res_vy.status_code == 200:
        vy_data = res_vy.json()
        assert vy_data["total"] >= 1
        assert len(vy_data["rows"]) == vy_data["total"]
        # Kiểm tra không có dấu '+' trong TD_soThuTuThua
        for r in vy_data["rows"]:
            assert "+" not in str(r.get("TD_soThuTuThua", ""))
        export_rows = vy_data["rows"][:5]
    else:
        assert res_vy.status_code in [200, 404]
        export_rows = [{"STT": 1, "GCN_soPhatHanh": "CH 123456", "TD_soThuTuThua": "12"}]

    # 3. Test GET /chuyen-doi/load-from-markdown-db
    res_mdb = client.get("/chuyen-doi/load-from-markdown-db")
    assert res_mdb.status_code == 200, res_mdb.text
    mdb_data = res_mdb.json()
    assert "total" in mdb_data
    assert "rows" in mdb_data
    assert isinstance(mdb_data["rows"], list)

    # 4. Test GET /api/v1/raw-ocr/to-129-rows
    res_r129 = client.get("/api/v1/raw-ocr/to-129-rows")
    assert res_r129.status_code == 200, res_r129.text
    r129_data = res_r129.json()
    assert "total" in r129_data
    assert "rows" in r129_data

    # 5. Test POST /chuyen-doi/export
    res_exp = client.post("/chuyen-doi/export", json={
        "rows": export_rows,
        "filename": "Test_Export_5_Rows.xlsx"
    })
    assert res_exp.status_code == 200
    assert len(res_exp.content) > 1000


if __name__ == "__main__":
    print("Running test_columns_json...")
    test_columns_json()
    print("PASSED test_columns_json")

    print("Running test_mapper_and_exporter...")
    test_mapper_and_exporter()
    print("PASSED test_mapper_and_exporter")

    print("Running test_map_merged_to_rows_multi_parcel...")
    test_map_merged_to_rows_multi_parcel()
    print("PASSED test_map_merged_to_rows_multi_parcel")

    print("Running test_api_endpoints...")
    test_api_endpoints()
    print("PASSED test_api_endpoints")

    print("\nALL TESTS PASSED SUCCESSFULLY! 100% OK")
