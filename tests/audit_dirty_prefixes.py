import sys
import io
import re
from collections import Counter
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook('KetQua_ChuyenDoi_129Cot.xlsx', data_only=True)
ws = wb.active

pattern_prefixes = Counter()
full_list = []

for r in range(5, ws.max_row + 1):
    v = str(ws.cell(row=r, column=30).value or '').strip()
    if not v or v == 'None':
        continue
    
    m = re.search(r'^(.*?)\b(Th[ôoóòõọỏơớờỡợởôốồỗộổaáàãạả]n|Xóm|Bản|Tổ|Đồng|Khu)\b', v, re.IGNORECASE)
    if m:
        prefix = m.group(1).strip(' -:;,./')
        if prefix:
            pattern_prefixes[prefix] += 1
            full_list.append((r, prefix, v))

print(f'Total rows with prefixes before Thon/Xom/Ban: {len(full_list)}')
print('\nUnique prefix patterns:')
for p, count in pattern_prefixes.most_common():
    print(f'  {count} rows : "{p}"')

print('\nSample rows:')
for r, p, v in full_list[:20]:
    print(f'  Row {r} [{p}] -> {v}')
