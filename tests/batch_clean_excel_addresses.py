import sys
import io
import os
import openpyxl

sys.path.insert(0, os.path.abspath('ocr-so-do/backend/src'))
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from ocr_so_do.application.projections.cadastral_129_mapper import Cadastral129Mapper

excel_path = 'KetQua_ChuyenDoi_129Cot.xlsx'
wb = openpyxl.load_workbook(excel_path)
ws = wb.active

col_map = {}
for c in range(1, ws.max_column + 1):
    val = ws.cell(row=1, column=c).value
    if val:
        col_map[str(val).strip()] = c

print(f'Total columns: {len(col_map)}')

updated_chu_addr = 0
updated_vc_addr = 0
updated_td_addr = 0

for r in range(5, ws.max_row + 1):
    # 1. Update CHU Address
    chu_dc = str(ws.cell(row=r, column=col_map['CHU_diaChiChiTiet']).value or '').strip()
    if chu_dc and chu_dc != 'None':
        clean_chu = Cadastral129Mapper.clean_address(chu_dc)
        if clean_chu != chu_dc:
            dec = Cadastral129Mapper.decompose_address(clean_chu)
            ws.cell(row=r, column=col_map['CHU_diaChiChiTiet'], value=dec['dia_chi_chi_tiet'])
            ws.cell(row=r, column=col_map['CHU_soNha'], value=dec['so_nha'])
            ws.cell(row=r, column=col_map['CHU_tenDuongPho'], value=dec['ten_duong_pho'])
            ws.cell(row=r, column=col_map['CHU_tenTDP'], value=dec['ten_tdp'])
            ws.cell(row=r, column=col_map['CHU_tenXa'], value=dec['ten_xa'])
            ws.cell(row=r, column=col_map['CHU_tenHuyen'], value=dec['ten_huyen'])
            ws.cell(row=r, column=col_map['CHU_tenTinh'], value=dec['ten_tinh'])
            updated_chu_addr += 1

    # 2. Update VC Address
    vc_dc = str(ws.cell(row=r, column=col_map['VC_diaChiChiTiet']).value or '').strip()
    if vc_dc and vc_dc != 'None':
        clean_vc = Cadastral129Mapper.clean_address(vc_dc)
        if clean_vc != vc_dc:
            dec = Cadastral129Mapper.decompose_address(clean_vc)
            ws.cell(row=r, column=col_map['VC_diaChiChiTiet'], value=dec['dia_chi_chi_tiet'])
            ws.cell(row=r, column=col_map['VC_soNha'], value=dec['so_nha'])
            ws.cell(row=r, column=col_map['VC_tenDuongPho'], value=dec['ten_duong_pho'])
            ws.cell(row=r, column=col_map['VC_tenTDP'], value=dec['ten_tdp'])
            ws.cell(row=r, column=col_map['VC_tenXa'], value=dec['ten_xa'])
            ws.cell(row=r, column=col_map['VC_tenHuyen'], value=dec['ten_huyen'])
            ws.cell(row=r, column=col_map['VC_tenTinh'], value=dec['ten_tinh'])
            updated_vc_addr += 1

    # 3. Update TD Address
    td_dc = str(ws.cell(row=r, column=col_map['TD_diaChiChiTiet']).value or '').strip()
    if td_dc and td_dc != 'None':
        clean_td = Cadastral129Mapper.clean_address(td_dc)
        if clean_td != td_dc:
            dec = Cadastral129Mapper.decompose_address(clean_td)
            ws.cell(row=r, column=col_map['TD_diaChiChiTiet'], value=dec['dia_chi_chi_tiet'])
            if 'TD_soNha' in col_map: ws.cell(row=r, column=col_map['TD_soNha'], value=dec['so_nha'])
            if 'TD_tenDuongPho' in col_map: ws.cell(row=r, column=col_map['TD_tenDuongPho'], value=dec['ten_duong_pho'])
            if 'TD_tenTDP' in col_map: ws.cell(row=r, column=col_map['TD_tenTDP'], value=dec['ten_tdp'])
            if 'TD_tenXa' in col_map: ws.cell(row=r, column=col_map['TD_tenXa'], value=dec['ten_xa'])
            if 'TD_tenHuyen' in col_map: ws.cell(row=r, column=col_map['TD_tenHuyen'], value=dec['ten_huyen'])
            if 'TD_tenTinh' in col_map: ws.cell(row=r, column=col_map['TD_tenTinh'], value=dec['ten_tinh'])
            updated_td_addr += 1

print(f'Rows updated in CHU_diaChiChiTiet: {updated_chu_addr}')
print(f'Rows updated in VC_diaChiChiTiet : {updated_vc_addr}')
print(f'Rows updated in TD_diaChiChiTiet : {updated_td_addr}')

wb.save(excel_path)
print('Successfully saved updated KetQua_ChuyenDoi_129Cot.xlsx!')
