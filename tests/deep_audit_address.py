import sys
import io
import re
from collections import Counter
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook('KetQua_ChuyenDoi_129Cot.xlsx', data_only=True)
ws = wb.active

for col_idx in [30, 48, 77]:
    col_name = str(ws.cell(row=1, column=col_idx).value or '').strip()
    print(f"\n==========================================")
    print(f"AUDITING COL {col_idx}: {col_name}")
    print(f"==========================================")
    
    total_dirty = 0
    categories = Counter()
    samples = []
    
    for r in range(5, ws.max_row + 1):
        val = str(ws.cell(row=r, column=col_idx).value or '').strip()
        if not val or val == 'None':
            continue
            
        val_l = val.lower()
        issue = None
        
        # 1. Prefix garbage (OCR of 'Địa chỉ thường trú')
        # Check if text precedes the administrative entity
        m_admin = re.search(r'\b(Th[ôoóòõọỏơớờỡợởôốồỗộổaáàãạả]n|Xóm|Bản|Tổ|Đồng|Khu|Số\s*\d+|Đội|Đoàn)\b', val, re.IGNORECASE)
        if m_admin and m_admin.start() > 0:
            prefix = val[:m_admin.start()].strip(' -:;,./')
            if prefix:
                issue = f"Tiền tố OCR rác đứng trước thôn/xóm: '{prefix}'"
                categories["Tiền tố OCR rác đứng trước thôn/xóm"] += 1
        
        # 2. Typos of 'Thôn' (Thòa, Thòn, Thòm, Thỏa, Thôa, Thơn, Thơ, Thỏo...)
        if not issue:
            m_thon_typo = re.match(r'^(Th[òóỏõọôốồổỗộơớờởỡợaáàảãạ]m|Th[òóỏõọôốồổỗộơớờởỡợaáàảãạ]n|Th[òóỏõọôốồổỗộơớờởỡợaáàảãạ]a|Th[òóỏõọôốồổỗộơớờởỡợaáàảãạ]o|Thơ)\s+', val, re.IGNORECASE)
            if m_thon_typo and m_thon_typo.group(1).lower() not in ['thôn', 'thon']:
                issue = f"Sai chính tả từ 'Thôn' -> '{m_thon_typo.group(1)}'"
                categories["Sai chính tả từ 'Thôn'"] += 1
                
        # 3. Suffix personal leakage (spouse, year of birth, CMND)
        if not issue:
            has_suffix = any(k in val_l for k in ['và ba', 'va ba', 'và bà', 'va bà', 'và ông', 'va ong', 'vợ là', 'chồng là', 'sinh năm', 'năm sinh', 'cmnd', 'cccd'])
            if has_suffix:
                issue = "Dính thông tin nhân thân/người đồng sở hữu ở đuôi"
                categories["Dính thông tin nhân thân ở đuôi"] += 1

        # 4. Leading symbols or single chars
        if not issue:
            if re.match(r'^[^\w\s]*[a-zA-Z0-9]{1,3}\s+', val) and not any(val_l.startswith(k) for k in ['số ', 'tổ ']):
                issue = "Dính ký tự rác đơn lẻ ở đầu"
                categories["Dính ký tự rác đơn lẻ ở đầu"] += 1
                
        if issue:
            total_dirty += 1
            if len(samples) < 15:
                samples.append((r, issue, val))
                
    print(f"Total problematic rows found: {total_dirty}")
    for cat, count in categories.items():
        print(f"  - {cat}: {count} dòng")
    print("\nSample rows:")
    for r, issue, val in samples:
        print(f"  Row {r}: [{issue}] -> {val}")
